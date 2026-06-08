# spell_checker.py - Corrector ortografico para JARVIS Mark XXXIX
# Archivo: C:\Users\jose1\Mark-XXXIX\actions\spell_checker.py

import sys
import time
import subprocess
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR / "actions"))
from gemini_helper import ask_gemini


def _encontrar_archivo_word(nombre_hint=None):
    """Encuentra el archivo Word mas reciente."""
    lugares = [
        Path("C:\\Users\\jose1\\Desktop"),
        Path("C:\\Users\\jose1\\Documents"),
        Path("C:\\Users\\jose1\\Downloads"),
    ]
    archivos = []
    for lugar in lugares:
        if lugar.exists():
            archivos.extend(lugar.glob("*.docx"))
            archivos.extend(lugar.glob("*.doc"))

    if not archivos:
        return None

    if nombre_hint:
        nombre_hint = nombre_hint.lower()
        for f in archivos:
            if nombre_hint in f.name.lower():
                return f

    return max(archivos, key=lambda f: f.stat().st_mtime)


def _corregir_word(ruta_archivo, player=None):
    """Corrige ortografia en Word mandando TODO el texto de una sola vez a la IA."""
    try:
        from docx import Document

        doc = Document(str(ruta_archivo))

        # Extraer todo el texto junto con indices de parrafos
        parrafos_info = []
        for i, parrafo in enumerate(doc.paragraphs):
            texto = parrafo.text.strip()
            if texto and len(texto) > 3:
                parrafos_info.append((i, texto))

        if not parrafos_info:
            return 0, 0, None

        # Construir un solo texto con separadores para mandarlo de una vez
        separador = "\n---PARRAFO---\n"
        texto_completo = separador.join([t for _, t in parrafos_info])

        if player:
            player.write_log(f"[SpellChecker] Corrigiendo {len(parrafos_info)} parrafos en una sola llamada...")

        # UNA SOLA llamada a la IA
        prompt = f"""Eres un corrector ortografico experto en español.

TAREA: Corrige SOLO los errores ortograficos del texto. Los parrafos estan separados por "---PARRAFO---".

REGLAS:
- Corrige SOLO errores de ortografia (letras incorrectas, acentos faltantes)
- NO cambies palabras, estilo ni significado
- Mantén EXACTAMENTE los separadores "---PARRAFO---" entre parrafos
- Retorna SOLO el texto corregido con los mismos separadores

TEXTO:
{texto_completo}"""

        resultado = ask_gemini(prompt)

        if not resultado:
            return 0, len(parrafos_info), None

        # Separar los parrafos corregidos
        parrafos_corregidos = resultado.strip().split("---PARRAFO---")
        parrafos_corregidos = [p.strip() for p in parrafos_corregidos]

        # Aplicar correcciones al documento
        cambios = 0
        for idx, (i_parrafo, texto_original) in enumerate(parrafos_info):
            if idx < len(parrafos_corregidos):
                texto_nuevo = parrafos_corregidos[idx]
                if texto_nuevo and texto_nuevo != texto_original:
                    parrafo = doc.paragraphs[i_parrafo]
                    # Guardar formato
                    if parrafo.runs:
                        bold      = parrafo.runs[0].bold
                        italic    = parrafo.runs[0].italic
                        font_size = parrafo.runs[0].font.size
                        font_name = parrafo.runs[0].font.name
                        # Limpiar y escribir
                        for run in parrafo.runs:
                            run.text = ""
                        parrafo.runs[0].text   = texto_nuevo
                        parrafo.runs[0].bold   = bold
                        parrafo.runs[0].italic = italic
                        if font_size:
                            parrafo.runs[0].font.size = font_size
                        if font_name:
                            parrafo.runs[0].font.name = font_name
                    cambios += 1

        # Guardar
        ruta_corregido = ruta_archivo.parent / (ruta_archivo.stem + "_corregido.docx")
        doc.save(str(ruta_corregido))
        subprocess.Popen(["explorer", str(ruta_corregido)])
        time.sleep(1)

        return cambios, len(parrafos_info), str(ruta_corregido)

    except Exception as e:
        print(f"[SpellChecker] Error: {e}")
        import traceback
        traceback.print_exc()
        return 0, 0, None


def spell_checker(parameters, player=None, speak=None):
    archivo_param = parameters.get("archivo", "").strip()
    nombre_param  = parameters.get("nombre", "").strip()

    if speak:
        speak("Buscando el documento, jefe.")

    # Encontrar archivo
    ruta = None
    if archivo_param:
        ruta = Path(archivo_param)
        if not ruta.exists():
            return f"No encontré '{archivo_param}', jefe."
    elif nombre_param:
        ruta = _encontrar_archivo_word(nombre_param)
        if not ruta:
            return f"No encontré ningún archivo llamado '{nombre_param}', jefe."
    else:
        ruta = _encontrar_archivo_word()
        if not ruta:
            return "No encontré documentos Word, jefe. Dígame el nombre del archivo."

    if player:
        player.write_log(f"[SpellChecker] Archivo: {ruta.name}")

    if speak:
        speak(f"Corrigiendo '{ruta.name}', jefe. Un momento.")

    extension = ruta.suffix.lower()
    if extension in (".docx", ".doc"):
        cambios, total, ruta_corregido = _corregir_word(ruta, player)
    else:
        return f"Formato '{extension}' no soportado. Funciona con .docx"

    if not ruta_corregido:
        return "Tuve un problema al corregir el documento, jefe."

    if speak:
        if cambios > 0:
            speak(
                f"Listo jefe. Corregí {cambios} párrafos en '{ruta.name}'. "
                f"Guardé el archivo como '{Path(ruta_corregido).name}' y lo abrí."
            )
        else:
            speak(f"Revisé '{ruta.name}' y no encontré errores ortográficos, jefe.")

    return f"Corrección: {cambios}/{total} párrafos corregidos. Archivo: {ruta_corregido}"
